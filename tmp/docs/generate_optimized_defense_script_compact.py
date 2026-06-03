from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


OUT_DIR = Path("output/doc")
DOCX_OUT = OUT_DIR / "答辩讲稿_优化版.docx"
MD_OUT = OUT_DIR / "答辩讲稿_优化版.md"


slides = [
    ("第1页 标题页", "约25秒", "各位老师好，我是22级自动化1班的傅懋杰。我的毕业设计题目是《面向设备健康状态的智能作业车间调度系统设计与实现》。接下来我会按照问题背景、数据与文献基础、系统设计、实验结果和自我评价这几个部分进行汇报。"),
    ("第2页 柔性作业车间调度问题", "约35秒", "本文的基础问题是柔性作业车间调度问题，也就是FJSP。简单来说，它不只要决定工序先后顺序，还要决定每道工序分配到哪一台可选机器上。图中的析取图表达了两类约束：同一作业内工序有先后关系，同一机器上的工序又会竞争加工资源。"),
    ("第3页 FJSP的现实应用场景", "约30秒", "FJSP在机械加工、半导体制造、汽车零部件生产，以及企业内小批量、多品种生产中都很常见。这些场景的共同特点是订单类型多、机器资源有限、加工路径不唯一，所以调度结果会直接影响交期、设备利用率和生产稳定性。"),
    ("第4页 RUL概念引入", "约40秒", "传统调度通常默认机器只分为空闲和忙碌，但真实设备会随着加工逐渐退化。RUL，也就是剩余使用寿命，用来描述设备距离失效还剩多少可用寿命。引入RUL以后，调度系统不仅能看到机器是否空闲，还能看到机器是否健康，从而在故障发生前安排维护，减少被动停机和过早维护之间的矛盾。"),
    ("第5页 课题任务、目的与意义", "约45秒", "因此，本文的任务是构建一个同时考虑设备退化、RUL预测、维护动作和动态FJSP调度的仿真与决策框架，并实现RUL预测、维护决策和调度决策三个模块。本文的目的，是把RUL健康信号引入生产调度和维护决策闭环，在动态作业到达、交期压力变化和设备退化并存的条件下，降低总成本、拖期成本和故障风险。它的意义主要在于：更贴近真实车间的生产-退化-维护过程，并通过不同算法组合比较调度器和维护器之间的适配关系。"),
    ("第6页 原始数据集来源", "约35秒", "本文使用的数据集是Kaggle上的Preventive to Predictive Maintenance dataset。该数据来自过滤测试台实验，粉尘进入滤材后，滤材逐渐堵塞，两端压差随时间上升，达到阈值后可视为失效。虽然它不是机床数据，但它提供了清晰的退化过程和RUL标签，适合本文构造设备健康状态。"),
    ("第7页 数据字段与本文使用方式", "约40秒", "数据集中训练集和测试集各有50条寿命序列，规模都接近四万行。Data_No表示寿命序列编号，Differential_pressure表示滤材两端压差，是本文RUL预测的核心输入；Dust_feed和Dust表示不同加尘工况；RUL是剩余寿命标签。本文主要使用压差序列，因为它与失效阈值直接相关，便于转化为0到1之间的健康状态。"),
    ("第8页 核心指导文献", "约45秒", "本文主要参考两篇文献。Luo等人在2021年的工作提供了动态FJSP、THDQN层级调度和规则动作设计思路；Ding等人在2025年的工作提供了将RUL预测与生产维护联合决策结合的思路，并使用GRU进行RUL预测。本文在此基础上做了改造：引入动态λ和DDT场景，加入多类维护算法，并系统比较THDQN、PPO与不同维护器的配对效果。"),
    ("第9页 总体系统设计", "约45秒", "本文系统分为三层：健康估计层、维护决策层和调度决策层。健康估计层用GRU输出归一化RUL；维护层根据健康状态和生产压力选择DN、IM或CM；调度层根据作业、机器、交期和健康状态选择派工规则。三者构成闭环：调度影响机器退化，维护改变机器可用性，RUL预测持续向维护和调度提供健康信息。"),
    ("第10页 RUL预测方法", "约55秒", "RUL模型输入是最近30个时间点的Differential_pressure。GRU读取这一段历史窗口，提取退化趋势，后接Dropout和MLP Head，最后通过Sigmoid输出0到1之间的归一化RUL。训练时选取Test_Data_CSV.csv中的Data_No=18，按Time升序划分为80%训练和20%验证。这里按时间切分而不是随机打乱，是为了保留退化序列的先后关系。需要说明的是，本文的RUL模型目标不是证明工业级泛化，而是在仿真中提供连续、方向正确的健康信号；在线使用时还做了单调修正，避免噪声导致局部“健康回升”。"),
    ("第11页 维护方法设计", "约35秒", "维护动作包含三类：DN表示不维护；IM表示非完全维护，可以理解为轻量保养或局部修复；CM表示恢复性维护，例如更换部件，使健康度大幅恢复。本文用维护时间、维护成本和健康恢复效果来描述维护动作，重点研究调度与维护如何协同决策，而不是建立具体设备的物理维修模型。"),
    ("第12页 三种维护算法", "约45秒", "维护算法设计了三种。Flat DDQN直接在DN、IM和CM中选动作，结构简单；Hier DDQN先判断是否维护，再判断选择IM还是CM，更接近“先判断要不要修，再判断怎么修”的实际逻辑；POMCP则使用粒子信念和蒙特卡洛树搜索做在线规划。三者分别代表单层学习、分层学习和规划式维护。"),
    ("第13页 调度方法设计", "约45秒", "调度部分比较THDQN和PPO。THDQN是层级结构，高层先选择优化目标，例如拖期、维护成本或二者折中，低层再在目标条件下选择具体规则。PPO作为扁平策略对照，直接学习从状态到6条规则的概率分布。这样设置是为了比较层级调度和扁平调度在健康约束下的适配能力。"),
    ("第14页 六条调度规则", "约40秒", "这六条规则可以理解为一组从效率优先到交期和风险优先的规则谱系。SRPT和SPT更偏向短任务和局部效率；EDD和Minimum Slack更关注交期压力；最大加权拖期风险规则更关注高损失订单。本文不是人工固定使用某一条规则，而是把这些可解释规则作为动作集合，让调度器学习何时切换。"),
    ("第15页 实验设计", "约40秒", "实验从两个维度构造动态场景。λ表示作业到达间隔均值，λ越小负载越高，取20、40和60；DDT表示交期紧度，DDT越小交期越紧，取1.0、1.5和2.0。算法组合方面，调度器包括THDQN和PPO，维护器包括flat_ddqn、hier_ddqn和pomcp，因此形成6组配对实验。"),
    ("第16页 实验结果总表", "约55秒", "从结果表可以看到，表现最好的是THDQN加hier_ddqn，总成本为113.170，延期成本为79.170，逾期工序数为7，逾期比例为2.48%。第二名是PPO加hier_ddqn，第三名是PPO加pomcp。最差的是THDQN加flat_ddqn，主要问题是延期成本过高。这说明hier_ddqn在两类调度器下都比较稳定；同时也说明THDQN的优势不是孤立存在的，只有和合适的维护器结合时才会转化为系统收益。"),
    ("第17页 甘特图结果", "约35秒", "这一页是最佳组合THDQN加hier_ddqn的甘特图。彩色块表示加工任务，带斜线的块表示维护动作，背景颜色表示不同λ和DDT场景。可以看到，维护动作实际插入到机器时间线上，会占用机器可用时间，并影响后续调度，因此维护不能作为外部事件单独处理。"),
    ("第18页 RUL曲线与维护窗口", "约40秒", "这一页从健康状态解释结果。不同颜色表示6台机器的归一化RUL，灰色竖带表示维护窗口。曲线下降代表加工导致退化，向上跳升代表维护恢复。可以看到，该策略不是频繁把所有机器修到满健康，而是在健康下降到一定程度后选择性介入，从而兼顾故障风险、维护成本和延期成本。"),
    ("第19页 目标与规则选择分布", "约50秒", "最后一页展示THDQN在不同λ和DDT下的行为变化。左侧是高层目标选择，右侧是低层规则选择。可以看到，不同负载和交期压力下，THDQN并没有固定使用同一个目标或规则，而是会在拖期、维护和平衡目标之间切换，并带动低层规则变化。这说明本文最重要的结论不是某一条规则最优，而是层级调度和层级维护之间存在结构匹配。"),
    ("结束语与自我评价", "约40秒", "总体来看，本文完成了RUL预测、维护决策和动态FJSP调度的联合框架，实现了GRU、三类维护算法和两类调度算法，并完成了6组组合实验。结果表明，THDQN加hier_ddqn在总成本和延期成本上表现最好。我的自我评价是：课题任务基本完成，并在参考文献基础上做了场景扩展和算法配对分析；不足是RUL数据规模和变量仍有限，维护模型也较抽象，后续可结合真实设备数据进一步验证。以上是我的汇报，恳请各位老师批评指正。"),
]


def set_font(run, font_name: str) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def add_spacing(paragraph) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:line"), "360")
    spacing.set(qn("w:lineRule"), "auto")
    spacing.set(qn("w:after"), "120")


def build_docx() -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for section in doc.sections:
        section.top_margin = Pt(54)
        section.bottom_margin = Pt(54)
        section.left_margin = Pt(63)
        section.right_margin = Pt(63)

    title = doc.add_heading("答辩个人陈述讲稿（优化版）", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        set_font(run, "黑体")

    note = doc.add_paragraph("建议总时长：约13-14分钟。口吻：正式、清楚、自然，不按论文正文朗读。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_spacing(note)

    for title_text, timing, body in slides:
        h = doc.add_heading(f"{title_text}（{timing}）", level=2)
        for run in h.runs:
            set_font(run, "黑体")
        p = doc.add_paragraph(body)
        p.paragraph_format.first_line_indent = Pt(22)
        add_spacing(p)
        for run in p.runs:
            set_font(run, "宋体")
    doc.save(DOCX_OUT)


def build_markdown() -> None:
    lines = [
        "# 答辩个人陈述讲稿（优化版）",
        "",
        "> 建议总时长：约13-14分钟。口吻：正式、清楚、自然，不按论文正文朗读。",
        "",
    ]
    for title_text, timing, body in slides:
        lines.extend([f"## {title_text}（{timing}）", "", body, ""])
    MD_OUT.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_markdown()
    build_docx()
    print(DOCX_OUT)
    print(MD_OUT)


if __name__ == "__main__":
    main()
